"""从 DVtest_output.xlsx（信号定义）生成 config/signals.json。

signals.json 是上位机的中文显示与转换语义配置源：
- description: 信号中文描述
- conversion:  整数位/小数位合成物理值 (V = i + f/frac_div)
- status:      状态信号语义 (normal 值 + 值→文本映射)

需依赖 openpyxl（仅生成工具用）。用法:
    python tools/gen_signals.py <xlsx> <dbc> -o config/signals.json
"""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

import openpyxl

from data.dbc_parser import load_dbc


# HSD 电流通道名 -> 对应的状态引脚信号
HSD_STATUS_MAP = {
    "HSD1": "HSD_Status_P230",
    "HSD2": "HSD_Status_P231",
    "HSD3": "HSD_Status_P232",
    "HSD4": "HSD_Status_P233",
    "HSD5": "HSD_Status_P234",
    "HSD6": "HSD_Status_P235",
    "HSD7": "HSD_Status_P236",
}

# 电压/温度 _AI 通道小数位进位：固件 ADC_Value_process 打包 frac = raw*100 (0-99)
# (DVtest_260810 DV_CAN.c 实证), 非 DBC 位长 8bit 推断的 /256
AI_FRAC_DIV = 100

# 英文 VAL_ 文本 -> 中文显示 (260810 DBC 新增英文枚举)
EN_VAL_TEXT = {
    "normal": "正常",
    "abnormal": "异常",
    "OverLoad": "过载",
    "OverLoadJudge": "过载判定",
    "OpenLoad": "开路",
    "ShortToVs": "对VS短路",
    "Error": "错误",
    "start": "启动",
    "running": "运行",
    "changechannel": "换通道",
    "Standby": "待机",
    "关闭": "关闭",
    "开启": "开启",
}

# 温度型整数/小数通道（非电压）
TEMP_UNITS = {"SOC_TEMP": "℃", "TEMP5152": "℃"}


def parse_xlsx(path: str):
    """返回 (msgs, signals)：msgs[id]=(name,cycle,len)，signals[(name)]=dict。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    msgs, signals, order = {}, {}, []
    cur_id = None
    for row in ws.iter_rows(min_row=2, values_only=True):
        def g(i):
            if i >= len(row) or row[i] is None:
                return ""
            return str(row[i]).strip().replace("\n", " ")
        msg = g(0)
        sig = g(7)
        if msg and not sig:
            try:
                cur_id = int(g(3), 0)
            except ValueError:
                cur_id = None
            cyc = g(5)
            msgs[cur_id] = {"name": msg, "cycle": int(cyc) if cyc.isdigit() else 0,
                            "len": int(g(6)) if g(6).isdigit() else 8}
        elif sig and cur_id is not None:
            bl = g(13)
            signals[sig] = {
                "msg_id": cur_id,
                "msg": msgs[cur_id]["name"],
                "desc": g(8),
                "order": g(9),
                "start_byte": g(11),
                "bit_len": int(bl) if bl.isdigit() else 0,
                "note": g(22),
            }
            order.append(sig)
    return msgs, signals


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("dbc")
    ap.add_argument("-o", "--output", default="config/signals.json")
    args = ap.parse_args()

    msgs, signals = parse_xlsx(args.xlsx)
    db = load_dbc(args.dbc)

    descriptions = {s: v["desc"] for s, v in signals.items()}
    conversion, status = {}, {}
    seen_i = set()

    # 1) 整数位/小数位 合并通道
    for sig, meta in signals.items():
        ai_f = "_AI_i"
        adc_f = "_adc_i"
        if sig.endswith(ai_f) or sig.endswith(adc_f):
            is_ai = sig.endswith(ai_f)
            suffix = ai_f if is_ai else adc_f
            frac_suffix = "_AI_f" if is_ai else "_adc_f"
            base = sig[: -len(suffix)]
            frac = f"{base}{frac_suffix}"
            if frac not in signals or sig in seen_i:
                continue
            seen_i.add(sig)
            # 小数位除数为 DBC 中小数信号实际位长 (8bit->/256, 4bit->/16)；
            # 但电压/温度 _AI 通道固件实为 frac*100 (ADC_Value_process)，恒定 /100
            frac_len = db.messages[meta["msg_id"]].signals[frac].length
            frac_div = AI_FRAC_DIV if is_ai else (1 << frac_len)
            unit = "A" if not is_ai else TEMP_UNITS.get(base, "V")
            for tail in ("整数位", "小数位", "整数", "小数", "低位", "高位"):
                if meta["desc"].endswith(tail):
                    desc = meta["desc"][: -len(tail)]
                    break
            else:
                desc = meta["desc"]
            chan = {
                "desc": desc,
                "msg": meta["msg"],
                "int_sig": sig,
                "frac_sig": frac,
                "frac_div": frac_div,
                "unit": unit,
            }
            status_sig = None
            if is_ai:
                status_sig = f"{base}_Status"
                if status_sig not in signals:
                    status_sig = None
            else:
                status_sig = HSD_STATUS_MAP.get(base)
                if status_sig is not None and status_sig not in signals:
                    status_sig = None
            if status_sig:
                chan["status_sig"] = status_sig
            conversion[base] = chan

    # 2) 状态信号语义: 从 DBC VAL_ 推导 normal + 文本（DBC 用 GBK 解码干净）
    for msg in db.messages.values():
        for sig in msg.signals.values():
            if not sig.values:
                continue
            vals = sig.values
            normal = 1  # 默认 1=正常
            texts = sorted(vals.items())
            # 识别正常文本（中英文）：{0:正常,1:异常} -> normal 0 ; {0:异常,1:正常} -> normal 1
            for raw_val, text in texts[:2]:
                if text in ("正常", "normal"):
                    normal = raw_val
                    break
            translated = {k: EN_VAL_TEXT.get(v, v) for k, v in vals.items()}
            status[sig.name] = {
                "desc": signals.get(sig.name, {}).get("desc") or sig.comment or sig.name,
                "msg": msg.name,
                "normal": normal,
                "values": translated,
            }

    # 3) 无 VAL_ 的状态信号（供电状态、未枚举）默认 normal=1
    for sig, meta in signals.items():
        if sig.endswith("_Status") and sig not in status:
            status[sig] = {"desc": meta["desc"], "msg": meta["msg"],
                           "normal": 1, "values": {}}

    out = {
        "descriptions": dict(sorted(descriptions.items())),
        "conversion": dict(sorted(conversion.items())),
        "status": dict(sorted(status.items())),
    }
    Path(args.output).parent.mkdir(parents=True, exist_ok=True)
    Path(args.output).write_text(
        json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"[gen_signals] {len(descriptions)} signals, "
          f"{len(conversion)} merged channels, {len(status)} status "
          f"-> {args.output}")


if __name__ == "__main__":
    main()